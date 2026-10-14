"""
MektepApp AI — учебный AI-ассистент для студентов (Streamlit + Google Gemini).

Режимы:
  1. Конспектирование и Q&A — структурированный конспект с главными мыслями,
     вопросами для самопроверки, чатом по материалу и экспортом в Word (.docx).
  2. Генератор Excel-задач — решение расчётных задач с пошаговым объяснением
     и настоящим файлом Excel (.xlsx) с живыми формулами (openpyxl).
  3. Быстрое решение ДЗ — чат-репетитор: объяснения с примерами, фото задач.

Запуск:
  pip install -r requirements.txt
  streamlit run app.py            # или: python app.py

Ключ берётся из файла .env (GEMINI_API_KEY=...) или вводится в боковой панели.
"""

import hashlib
import html
import io
import logging
import math
import os
import re
import sys
import time
from collections.abc import Callable, Iterable, Iterator
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Literal

import httpx
import pandas as pd
import streamlit as st
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from dotenv import load_dotenv, set_key
from google import genai
from google.genai import errors as genai_errors
from google.genai import types
from openpyxl import Workbook
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import FORMULAE
from pydantic import BaseModel, Field, ValidationError

log = logging.getLogger("mektepapp")

# ════════════════════════════════════════════════════════════════════════════
# Конфигурация
# ════════════════════════════════════════════════════════════════════════════

APP_NAME = "MektepApp AI"
APP_DIR = Path(__file__).resolve().parent
ENV_PATH = APP_DIR / ".env"

load_dotenv(ENV_PATH)
load_dotenv(Path.cwd() / ".env")

DEFAULT_MODEL = os.getenv("GEMINI_MODEL", "").strip() or "gemini-2.5-flash"
MODEL_OPTIONS = list(
    dict.fromkeys([DEFAULT_MODEL, "gemini-2.5-flash", "gemini-flash-latest", "gemini-2.5-flash-lite"])
)
API_KEY_URL = "https://aistudio.google.com/apikey"

REQUEST_TIMEOUT_MS = 180_000  # тайм-аут одного запроса к Gemini
MAX_FILE_MB = 10  # лимит для загружаемых файлов (PDF/DOCX/фото)
MAX_TEXT_CHARS = 200_000  # лимит для вставленного текста
MAX_CONTEXT_CHARS = 60_000  # сколько исходного текста передаём в чат по материалу

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

LANGUAGES = {
    "Русский": "на русском языке",
    "Қазақша": "на казахском языке (қазақ тілінде)",
    "English": "на английском языке (in English)",
}
LEVELS = ["Школа", "Колледж", "Университет"]
SUBJECTS = [
    "Любой предмет", "Математика", "Физика", "Химия", "Биология",
    "Информатика и программирование", "Экономика и бухучёт", "История",
    "География", "Литература", "Русский язык", "Казахский язык",
    "Английский язык", "Право", "Философия",
]

NOTES_SOURCES = ["Тема", "Текст", "Файл"]
NOTES_DETAIL = {
    "Кратко": "3–4 раздела по 2–4 тезиса, 3–4 главные мысли, 3–6 терминов",
    "Стандартно": "4–6 разделов по 3–6 тезисов, 4–6 главных мыслей, 5–10 терминов",
    "Подробно": "6–9 разделов по 4–8 тезисов с примерами и пояснениями, 5–8 главных мыслей, 8–15 терминов",
}
HW_STYLES = {
    "Кратко": "Отвечай кратко: суть и решение в 3–6 предложениях, без лишних деталей.",
    "Подробно": "Дай подробное объяснение: ответ, разбор, пример и типичные ошибки.",
    "Пошагово": "Объясняй максимально пошагово, как новичку: каждый шаг отдельно, "
                "с промежуточными результатами и проверкой в конце.",
}

EXCEL_EXAMPLES = [
    ("💰 Зарплата", "Рассчитай заработную плату 5 сотрудников за месяц. Оклады: 250 000, 320 000, "
                   "180 000, 410 000 и 275 000 ₸. Удержи пенсионные взносы 10% от оклада и подоходный "
                   "налог 10% от суммы после взносов. Найди сумму к выдаче по каждому и итоги."),
    ("🏦 Кредит", "Составь график погашения аннуитетного кредита: 1 500 000 ₸ на 12 месяцев под 18% "
                 "годовых. По месяцам покажи платёж, проценты, основной долг и остаток. "
                 "Посчитай переплату."),
    ("📦 Учёт ФИФО", "Учёт товаров методом ФИФО: начальный остаток 100 шт. по 500 ₸, закупки 200 шт. "
                    "по 550 ₸ и 150 шт. по 600 ₸, продано 320 шт. Найди себестоимость продаж "
                    "и стоимость конечного остатка."),
    ("📈 Рентабельность", "Магазин продал 4 товара. Выручка: 1 200 000, 850 000, 430 000, 2 100 000 ₸; "
                         "себестоимость: 900 000, 610 000, 350 000, 1 650 000 ₸. Найди валовую прибыль "
                         "и рентабельность продаж по каждому товару и в целом."),
    ("🎓 Оценки", "Придумай оценки 8 студентов по 4 предметам (по 100-балльной шкале). Посчитай "
                 "средний балл каждого студента, средний балл по каждому предмету и определи лучшего студента."),
]

HW_STARTERS = [
    ("📐", "Реши уравнение 2x² − 7x + 3 = 0 и объясни каждый шаг"),
    ("⚡", "Объясни закон Ома простыми словами и приведи пример из жизни"),
    ("🧪", "Уравняй реакцию Al + O₂ → Al₂O₃ и объясни, как это делать"),
    ("💻", "Что такое рекурсия? Покажи понятный пример на Python"),
    ("📊", "Чем отличается дебет от кредита? Приведи пример проводки"),
    ("✍️", "Помоги составить план эссе «Роль технологий в образовании»"),
]
NOTES_STARTERS = [
    ("💡", "Объясни главную идею простыми словами"),
    ("🌍", "Приведи пример из реальной жизни"),
    ("🎯", "Какие вопросы по этой теме чаще всего бывают на экзамене?"),
    ("🗣️", "Составь план устного ответа на 2 минуты"),
]

CHAT_FILE_TYPES = ["png", "jpg", "jpeg", "webp", "pdf", "txt"]
MIME_BY_EXT = {
    "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "webp": "image/webp",
    "pdf": "application/pdf", "txt": "text/plain", "md": "text/plain",
}
USER_AVATAR = ":material/person:"
AI_AVATAR = ":material/auto_awesome:"

# Подписи внутри документов Word/Excel — на выбранном языке ответов.
DOC_LABELS = {
    "Русский": {
        "summary": "Кратко о главном", "key_ideas": "Главные мысли", "sections": "Конспект",
        "terms": "Ключевые термины", "term": "Термин", "definition": "Определение",
        "qa": "Вопросы для самопроверки", "answer": "Ответ", "conclusion": "Вывод",
        "solution": "Решение", "task": "Условие", "steps": "Пошаговое решение",
        "formulas": "Ключевые формулы", "formula": "Формула", "meaning": "Что считает",
        "created": "Создано в MektepApp AI", "page": "Стр.",
    },
    "Қазақша": {
        "summary": "Қысқаша мазмұны", "key_ideas": "Негізгі ойлар", "sections": "Конспект",
        "terms": "Негізгі терминдер", "term": "Термин", "definition": "Анықтамасы",
        "qa": "Өзін-өзі тексеру сұрақтары", "answer": "Жауабы", "conclusion": "Қорытынды",
        "solution": "Шешімі", "task": "Есептің шарты", "steps": "Қадамдық шешім",
        "formulas": "Негізгі формулалар", "formula": "Формула", "meaning": "Түсіндірме",
        "created": "MektepApp AI көмегімен жасалды", "page": "Бет",
    },
    "English": {
        "summary": "Summary", "key_ideas": "Key ideas", "sections": "Notes",
        "terms": "Key terms", "term": "Term", "definition": "Definition",
        "qa": "Self-check questions", "answer": "Answer", "conclusion": "Conclusion",
        "solution": "Solution", "task": "Task", "steps": "Step-by-step solution",
        "formulas": "Key formulas", "formula": "Formula", "meaning": "What it does",
        "created": "Created with MektepApp AI", "page": "Page",
    },
}


@dataclass(frozen=True)
class Mode:
    title: str
    caption: str
    subtitle: str
    emoji: str
    icon: str


MODES = {
    "notes": Mode(
        "Конспектирование и Q&A", "Конспект, вопросы и Word",
        "Тема, текст или файл → структурированный конспект с главными мыслями, "
        "вопросами для самопроверки и готовым документом Word.",
        "📝", ":material/edit_note:",
    ),
    "excel": Mode(
        "Генератор Excel-задач", "Расчёты и таблицы в .xlsx",
        "Расчёты, таблицы и учёт: пошаговое объяснение решения и настоящий файл Excel "
        "с живыми формулами.",
        "📊", ":material/table_view:",
    ),
    "homework": Mode(
        "Быстрое решение ДЗ", "Объяснения с примерами",
        "Задайте вопрос по любому предмету или прикрепите фото задачи — "
        "получите понятный разбор с примерами.",
        "🎓", ":material/school:",
    ),
}


@dataclass
class Ctx:
    """Настройки текущего запуска, общие для всех режимов."""

    client: genai.Client | None
    model: str
    language: str
    lang_name: str
    level: str
    deep: bool

    @property
    def ready(self) -> bool:
        return self.client is not None

    @property
    def labels(self) -> dict[str, str]:
        return DOC_LABELS[self.lang_name]


# ════════════════════════════════════════════════════════════════════════════
# Схемы структурированных ответов Gemini
# ════════════════════════════════════════════════════════════════════════════

class KeyIdea(BaseModel):
    title: str = Field(description="Короткий заголовок мысли, 3–7 слов")
    text: str = Field(description="Пояснение в 1–2 предложениях")


class NoteSection(BaseModel):
    heading: str = Field(description="Заголовок раздела")
    points: list[str] = Field(description="Тезисы раздела; каждый — законченная мысль")


class Term(BaseModel):
    term: str
    definition: str


class QAPair(BaseModel):
    question: str
    answer: str


class StudyNotes(BaseModel):
    title: str = Field(description="Название конспекта")
    summary: str = Field(description="Суть темы в 2–4 предложениях")
    key_ideas: list[KeyIdea] = Field(description="Главные мысли — самое важное для запоминания")
    sections: list[NoteSection] = Field(description="Конспект по разделам в логическом порядке")
    terms: list[Term] = Field(description="Ключевые термины с точными определениями")
    qa: list[QAPair] = Field(description="Вопросы для самопроверки с краткими ответами")
    conclusion: str = Field(description="Итоговый вывод в 2–3 предложениях")


class XlColumn(BaseModel):
    name: str = Field(description="Заголовок столбца; единицы измерения — в заголовке")
    kind: Literal["text", "integer", "number", "money", "percent"] = Field(description="Тип данных столбца")


class XlCell(BaseModel):
    value: str = Field(description="Значение ячейки; для формулы — уже рассчитанный результат")
    formula: str = Field(description="Формула Excel на английском, начинается с '='; пусто, если формулы нет")


class XlRow(BaseModel):
    cells: list[XlCell]


class XlSheet(BaseModel):
    name: str = Field(description="Имя листа, до 28 символов")
    columns: list[XlColumn]
    rows: list[XlRow] = Field(description="Строки данных; первая строка данных — строка 2 листа")


class XlFormulaNote(BaseModel):
    formula: str = Field(description="Формула Excel из таблицы")
    explanation: str = Field(description="Что считает формула, простыми словами")


class ExcelSolution(BaseModel):
    title: str = Field(description="Короткое название задачи")
    task: str = Field(description="Условие задачи в 1–3 предложениях")
    steps: list[str] = Field(description="Пошаговое решение с расчётами, 3–8 шагов")
    answer: str = Field(description="Итоговый ответ с конкретными числами")
    sheets: list[XlSheet] = Field(description="Листы Excel-файла (1–4)")
    key_formulas: list[XlFormulaNote] = Field(description="2–6 ключевых формул из таблиц")


# ════════════════════════════════════════════════════════════════════════════
# Промпты
# ════════════════════════════════════════════════════════════════════════════

def base_rules(ctx: Ctx) -> str:
    return (
        f"Ты — {APP_NAME}, доброжелательный и точный преподаватель-наставник. "
        f"Отвечай {ctx.language}. Уровень учащегося: {ctx.level.lower()} — подбирай сложность, "
        "термины и примеры под этот уровень. Не выдумывай факты, цифры и источники; "
        "если данных не хватает или ты не уверен — скажи об этом прямо."
    )


def notes_system(ctx: Ctx, detail: str, n_questions: int) -> str:
    return base_rules(ctx) + f"""

Задача: составить структурированный учебный конспект.
Правила:
• Опирайся на предоставленный материал. Если дана только тема — используй проверенные академические знания.
• key_ideas — самые важные мысли для запоминания: короткий заголовок и пояснение в 1–2 предложениях.
• sections — логичная структура темы; каждый тезис — законченная мысль. Ключевые термины выделяй **жирным**.
• terms — термины с точными определениями.
• qa — ровно {n_questions} вопросов для самопроверки (от простых к сложным) с краткими точными ответами.
• Объём: {NOTES_DETAIL[detail]}.
• Формулы пиши обычным текстом с Unicode-символами (x², √, π, ≤, →, Δ), без LaTeX и без символа $.
• Пиши ясно, без воды и повторов."""


def excel_system(ctx: Ctx) -> str:
    return base_rules(ctx) + """

Ты — эксперт по Excel, экономике, бухгалтерскому учёту, финансам и прикладной математике.
Задача: решить задачу студента и подготовить данные для настоящего файла Excel.

Таблицы (sheets):
• Один лист — одна таблица. Строка 1 — заголовки столбцов (columns), данные (rows) начинаются со строки 2, первый столбец — A.
• В каждой строке ровно столько ячеек (cells), сколько столбцов.
• Всё, что вычисляется, считай формулой Excel в поле formula: функции на английском (SUM, AVERAGE, ROUND, IF, SUMIF, SUMPRODUCT, PMT, IPMT, PPMT, VLOOKUP, INDEX, MATCH, MAX, MIN…), начинается с «=», аргументы через запятую, ссылки по раскладке выше (=B2*C2, =SUM(D2:D6), =ROUND(B2*0.1,2)). Ссылка на другой лист: ='Имя листа'!B2.
• Используй только классические функции Excel (без LET, LAMBDA, FILTER, XLOOKUP).
• value — всегда итоговое значение ячейки, для формулы — уже рассчитанный результат. Числа пиши без пробелов, валюты и единиц, десятичный разделитель — точка (12500.5); проценты — со знаком % (15%).
• Если формулы нет — formula = "".
• kind столбца: text — текст и коды, integer — целые числа, number — дробные, money — деньги, percent — проценты. Единицы и валюту указывай в заголовке: «Сумма, ₸».
• Исходные данные клади в отдельные ячейки, чтобы формулы на них ссылались. Добавляй строки «Итого» с формулами, где уместно.
• Имя листа — до 28 символов, без символов [ ] : * ? / \\. Не более 4 листов и 60 строк на листе.

Остальные поля:
• title — короткое название задачи; task — условие в 1–3 предложениях.
• steps — пошаговое решение простым языком (3–8 шагов) с расчётами.
• answer — итоговый ответ с конкретными числами.
• key_formulas — 2–6 ключевых формул из таблиц с пояснением.
Если данных не хватает — сделай разумные допущения и явно укажи их в steps."""


def homework_system(ctx: Ctx, subject: str, style: str) -> str:
    return base_rules(ctx) + f"""

Ты помогаешь с домашним заданием. Предмет: {subject}. Цель — чтобы ученик понял решение, а не просто списал ответ.
Структура ответа:
1. **Ответ** — коротко, 1–2 предложения (если вопрос предполагает ответ).
2. **Решение / объяснение** — по шагам, понятным языком.
3. **Пример** — похожий пример или аналогия из жизни.
4. **Проверь себя** — одно короткое задание для закрепления (если уместно).
{HW_STYLES[style]}
Оформление: Markdown, заголовки не крупнее ###. Математику пиши в LaTeX: $...$ внутри строки и $$...$$ отдельной строкой.
Символ $ используй только для формул (валюту пиши словами или знаком ₸).
Если прислано фото или файл с заданием — сначала кратко перепиши условие, затем решай.
Если в условии не хватает данных — укажи, какое допущение делаешь."""


def notes_chat_system(ctx: Ctx, result: dict) -> str:
    material = result["context"][:MAX_CONTEXT_CHARS] or "(исходный материал приложен в начале диалога)"
    return base_rules(ctx) + f"""

Отвечай на вопросы студента по учебному материалу и конспекту ниже.
• Опирайся на материал. Если ответа в нём нет — скажи об этом и ответь по общим знаниям, явно пометив это.
• Сначала короткий ответ, затем объяснение и пример. Markdown; формулы — LaTeX в $...$.

=== КОНСПЕКТ ===
{result["markdown"]}

=== ИСХОДНЫЙ МАТЕРИАЛ ===
{material}"""


# ════════════════════════════════════════════════════════════════════════════
# Gemini: клиент, запросы, ошибки
# ════════════════════════════════════════════════════════════════════════════

class AppError(Exception):
    """Ошибка с готовым понятным сообщением для пользователя."""


class Truncated(Exception):
    """Ответ оборван по лимиту длины (MAX_TOKENS)."""


BLOCKED_FINISH = {"SAFETY", "PROHIBITED_CONTENT", "BLOCKLIST", "SPII", "IMAGE_SAFETY"}


@st.cache_resource(show_spinner=False, max_entries=4)
def get_client(api_key: str) -> genai.Client:
    return genai.Client(
        api_key=api_key,
        http_options=types.HttpOptions(
            timeout=REQUEST_TIMEOUT_MS,
            retry_options=types.HttpRetryOptions(
                attempts=3, initial_delay=2.0, max_delay=12.0,
                http_status_codes=[500, 502, 503, 504],
            ),
        ),
    )


def thinking_config(model: str, deep: bool) -> types.ThinkingConfig | None:
    # «Размышления» можно отключить только у Gemini 2.5 Flash / Flash-Lite;
    # остальные модели работают в своём режиме по умолчанию.
    if not deep and model.startswith("gemini-2.5-flash"):
        return types.ThinkingConfig(thinking_budget=0)
    return None


def finish_name(response: types.GenerateContentResponse) -> str:
    if not response.candidates or response.candidates[0].finish_reason is None:
        return ""
    reason = response.candidates[0].finish_reason
    return getattr(reason, "name", str(reason))


def iter_text(ctx: Ctx, contents: list, config: types.GenerateContentConfig) -> Iterator[str]:
    """Потоковый ответ Gemini: текстовые фрагменты с проверкой блокировок и обрыва."""
    reason = ""
    for chunk in ctx.client.models.generate_content_stream(model=ctx.model, contents=contents, config=config):
        if chunk.prompt_feedback and chunk.prompt_feedback.block_reason:
            raise AppError("Запрос отклонён фильтрами безопасности Gemini. Переформулируйте его.")
        reason = finish_name(chunk) or reason
        if chunk.text:
            yield chunk.text
    if reason in BLOCKED_FINISH:
        raise AppError("Ответ остановлен фильтрами безопасности Gemini. Попробуйте переформулировать запрос.")
    if reason == "RECITATION":
        raise AppError("Ответ остановлен: он слишком близко цитирует защищённый источник. "
                       "Переформулируйте запрос.")
    if reason == "MAX_TOKENS":
        raise Truncated


def generate_json(ctx: Ctx, contents: list, *, system: str, schema: type[BaseModel], temperature: float,
                  on_progress: Callable[[int], None] | None = None) -> Any:
    """Ответ строго по JSON-схеме. Идёт потоком: виден прогресс, а долгая генерация не упирается в тайм-аут."""
    config = types.GenerateContentConfig(
        system_instruction=system,
        temperature=temperature,
        response_mime_type="application/json",
        response_schema=schema,
        thinking_config=thinking_config(ctx.model, ctx.deep),
    )
    text = ""
    try:
        for piece in iter_text(ctx, contents, config):
            text += piece
            if on_progress:
                on_progress(len(text))
    except Truncated:
        raise AppError("Ответ оказался слишком длинным и был обрезан. Сократите материал "
                       "или выберите детализацию «Кратко».") from None
    try:
        return schema.model_validate_json(text)
    except ValidationError:
        raise AppError("Gemini вернул ответ в неожиданном формате. Нажмите кнопку ещё раз.") from None


def stream_text(ctx: Ctx, contents: list, *, system: str, temperature: float = 0.5) -> Iterator[str]:
    """Потоковый текстовый ответ Gemini (для чатов)."""
    config = types.GenerateContentConfig(
        system_instruction=system,
        temperature=temperature,
        thinking_config=thinking_config(ctx.model, ctx.deep),
    )
    try:
        yield from iter_text(ctx, contents, config)
    except Truncated:
        yield "\n\n*…ответ обрезан из-за ограничения длины. Попросите продолжить.*"


def describe_error(exc: Exception, model: str) -> str:
    """Превращает любое исключение в понятное сообщение на русском."""
    if isinstance(exc, AppError):
        return str(exc)
    if isinstance(exc, genai_errors.APIError):
        code, message = exc.code or 0, exc.message or ""
        low = message.lower()
        if "api key not valid" in low or "api_key_invalid" in low:
            return "Неверный API-ключ. Проверьте ключ в боковой панели или в файле .env."
        if "location is not supported" in low:
            return "Gemini API недоступен в вашем регионе. Подключитесь через поддерживаемую сеть."
        if code == 429:
            delay = re.search(r"retry in ([\d.]+)s", message, re.I) or re.search(
                r"'retryDelay': '(\d+)s'", str(exc.details))
            wait = f" Повторите примерно через {math.ceil(float(delay.group(1)))} с." if delay else \
                " Подождите минуту и повторите."
            return "Превышен лимит запросов к Gemini (квота бесплатного тарифа)." + wait
        if code in (401, 403):
            return "Доступ запрещён: у ключа нет прав на Gemini API. Проверьте ключ в Google AI Studio."
        if code == 404:
            return (f"Модель «{model}» недоступна для этого ключа или снята с поддержки. "
                    "Выберите другую модель в боковой панели (например, gemini-flash-latest).")
        if code in (500, 502, 503, 504):
            return "Серверы Gemini сейчас перегружены или недоступны. Повторите попытку через минуту."
        return f"Ошибка Gemini API ({code}): {message or exc}"
    if isinstance(exc, httpx.TimeoutException):
        return "Gemini слишком долго не отвечает. Сократите запрос или повторите попытку."
    if isinstance(exc, httpx.HTTPError):
        return "Нет соединения с Gemini API. Проверьте интернет, VPN или прокси."
    log.exception("Непредвиденная ошибка", exc_info=exc)
    return f"Непредвиденная ошибка: {type(exc).__name__}: {exc}"


# ════════════════════════════════════════════════════════════════════════════
# Текстовые утилиты
# ════════════════════════════════════════════════════════════════════════════

_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_INLINE_MD = re.compile(r"(\*\*[^*]+?\*\*|`[^`]+`|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))")


def clean_text(value: Any) -> str:
    """Убирает управляющие символы, недопустимые в Word/Excel."""
    return _CONTROL_CHARS.sub("", str(value or "")).strip()


def esc(value: Any) -> str:
    return html.escape(clean_text(value))


def md_inline_html(text: str) -> str:
    """Безопасный HTML из строки с **жирным**, *курсивом* и `кодом`."""
    out = []
    for part in _INLINE_MD.split(clean_text(text)):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            out.append(f"<strong>{html.escape(part[2:-2])}</strong>")
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            out.append(f"<code>{html.escape(part[1:-1])}</code>")
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append(f"<em>{html.escape(part[1:-1])}</em>")
        else:
            out.append(html.escape(part))
    return "".join(out).replace("\n", "<br>")


def md_safe(text: str) -> str:
    """Текст для st.markdown без случайной LaTeX-разметки ($)."""
    return clean_text(text).replace("$", r"\$")


def slugify(text: str, fallback: str) -> str:
    slug = re.sub(r"[^\w\s-]", "", clean_text(text)).strip()
    return re.sub(r"[\s-]+", "_", slug)[:60] or fallback


def shorten(text: str, limit: int = 60) -> str:
    text = " ".join(clean_text(text).split())
    return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"


def read_source_file(upload) -> tuple[str, bytes | None]:
    """Читает загруженный файл: возвращает (текст, None) или ("", байты PDF)."""
    data = upload.getvalue()
    if len(data) > MAX_FILE_MB * 1024 * 1024:
        raise AppError(f"Файл больше {MAX_FILE_MB} МБ. Сократите документ или вставьте нужную часть текста.")
    name = upload.name.lower()
    if name.endswith(".pdf"):
        return "", data
    if name.endswith(".docx"):
        try:
            document = Document(io.BytesIO(data))
        except Exception as exc:
            raise AppError("Не удалось прочитать .docx — возможно, файл повреждён.") from exc
        lines = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(dict.fromkeys(cells)))
        text = "\n".join(lines)
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1251", errors="replace")
    if not text.strip():
        raise AppError("В файле не найден текст.")
    return text, None


# ════════════════════════════════════════════════════════════════════════════
# Экспорт конспекта: Markdown и Word (.docx)
# ════════════════════════════════════════════════════════════════════════════

INDIGO = RGBColor(0x43, 0x38, 0xCA)
VIOLET = RGBColor(0x6D, 0x28, 0xD9)
GRAY = RGBColor(0x6B, 0x72, 0x80)
INK = RGBColor(0x1F, 0x23, 0x2E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def notes_to_markdown(notes: StudyNotes, labels: dict) -> str:
    lines = [f"# {notes.title}", "", notes.summary, "", f"## {labels['key_ideas']}"]
    lines += [f"{i}. **{idea.title}** — {idea.text}" for i, idea in enumerate(notes.key_ideas, 1)]
    lines += ["", f"## {labels['sections']}"]
    for section in notes.sections:
        lines += ["", f"### {section.heading}", *[f"- {point}" for point in section.points]]
    if notes.terms:
        lines += ["", f"## {labels['terms']}", *[f"- **{t.term}** — {t.definition}" for t in notes.terms]]
    if notes.qa:
        lines += ["", f"## {labels['qa']}"]
        for i, item in enumerate(notes.qa, 1):
            lines += ["", f"**{i}. {item.question}**", "", f"{labels['answer']}: {item.answer}"]
    if notes.conclusion:
        lines += ["", f"## {labels['conclusion']}", "", notes.conclusion]
    return "\n".join(lines) + "\n"


def add_inline(paragraph, text: str, *, bold: bool = False, color: RGBColor | None = None) -> None:
    """Добавляет текст в абзац Word, превращая **жирный**, *курсив* и `код` в форматирование."""
    for part in _INLINE_MD.split(clean_text(text)):
        if not part:
            continue
        run_bold, italic, mono = bold, False, False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, run_bold = part[2:-2], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part, mono = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, italic = part[1:-1], True
        run = paragraph.add_run(part)
        run.bold, run.italic = run_bold, italic
        if mono:
            run.font.name = "Consolas"
        if color is not None:
            run.font.color.rgb = color


def style_cell(cell, fill: str, borders: dict[str, tuple[str, int, str]] | None = None) -> None:
    """Заливка и границы ячейки таблицы Word (порядок элементов важен для схемы OOXML)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    if borders is not None:
        tc_borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            val, size, color = borders.get(edge, ("nil", 0, "auto"))
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), val)
            if val != "nil":
                element.set(qn("w:sz"), str(size))
                element.set(qn("w:space"), "0")
                element.set(qn("w:color"), color)
            tc_borders.append(element)
        tc_pr.append(tc_borders)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def add_page_number(paragraph) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "6B7280")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    run_pr.append(color)
    run_pr.append(size)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(run_pr)
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def setup_document(document) -> None:
    for section in document.sections:
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2.2)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, color, size in (("Title", INDIGO, 24), ("Heading 1", INDIGO, 16), ("Heading 2", VIOLET, 13)):
        document.styles[name].font.color.rgb = color
        document.styles[name].font.size = Pt(size)


def build_notes_docx(notes: StudyNotes, labels: dict, source_label: str) -> bytes:
    document = Document()
    setup_document(document)
    document.core_properties.title = clean_text(notes.title)
    document.core_properties.author = APP_NAME
    document.core_properties.subject = clean_text(source_label)

    document.add_heading(clean_text(notes.title), level=0)
    meta = document.add_paragraph()
    run = meta.add_run(f"{labels['created']} · {datetime.now():%d.%m.%Y}")
    run.font.size, run.font.color.rgb = Pt(9), GRAY

    if notes.summary:
        document.add_heading(labels["summary"], level=1)
        add_inline(document.add_paragraph(), notes.summary)

    if notes.key_ideas:
        document.add_heading(labels["key_ideas"], level=1)
        table = document.add_table(rows=0, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        separator = ("single", 12, "FFFFFF")
        for i, idea in enumerate(notes.key_ideas, 1):
            cell = table.add_row().cells[0]
            style_cell(cell, "F3F0FF", {"left": ("single", 24, "7C3AED"), "top": separator, "bottom": separator})
            title = cell.paragraphs[0]
            title.paragraph_format.space_before = Pt(4)
            title.paragraph_format.space_after = Pt(2)
            add_inline(title, f"{i}. {idea.title}", bold=True, color=VIOLET)
            if idea.text:
                body = cell.add_paragraph()
                body.paragraph_format.space_after = Pt(4)
                add_inline(body, idea.text)

    if notes.sections:
        document.add_heading(labels["sections"], level=1)
        for section in notes.sections:
            document.add_heading(clean_text(section.heading), level=2)
            for point in section.points:
                add_inline(document.add_paragraph(style="List Bullet"), point)

    if notes.terms:
        document.add_heading(labels["terms"], level=1)
        table = document.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, (labels["term"], labels["definition"])):
            style_cell(cell, "4F46E5")
            add_inline(cell.paragraphs[0], text, bold=True, color=WHITE)
        header_pr = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        header_pr.append(repeat)
        for term in notes.terms:
            row = table.add_row().cells
            add_inline(row[0].paragraphs[0], term.term, bold=True)
            add_inline(row[1].paragraphs[0], term.definition)
        table.autofit = False
        for column, width in zip(table.columns, (Cm(4.6), Cm(12))):
            column.width = width
            for cell in column.cells:
                cell.width = width

    if notes.qa:
        document.add_heading(labels["qa"], level=1)
        for i, item in enumerate(notes.qa, 1):
            question = document.add_paragraph()
            question.paragraph_format.space_before = Pt(6)
            question.paragraph_format.keep_with_next = True
            add_inline(question, f"{i}. {item.question}", bold=True)
            answer = document.add_paragraph()
            answer.paragraph_format.left_indent = Cm(0.6)
            prefix = answer.add_run(f"{labels['answer']}: ")
            prefix.italic, prefix.font.color.rgb = True, GRAY
            add_inline(answer, item.answer)

    if notes.conclusion:
        document.add_heading(labels["conclusion"], level=1)
        add_inline(document.add_paragraph(), notes.conclusion)

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"{labels['created']}  ·  {labels['page']} ")
    run.font.size, run.font.color.rgb = Pt(9), GRAY
    add_page_number(footer)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ════════════════════════════════════════════════════════════════════════════
# Экспорт решения: Excel (.xlsx)
# ════════════════════════════════════════════════════════════════════════════

# Современные функции Excel хранятся в файле с префиксом _xlfn.
MODERN_FUNCTIONS = {
    "IFS", "MAXIFS", "MINIFS", "SWITCH", "TEXTJOIN", "CONCAT", "IFNA", "XLOOKUP", "XMATCH",
    "STDEV.S", "STDEV.P", "VAR.S", "VAR.P", "RANK.EQ", "RANK.AVG", "DAYS", "PERCENTILE.INC",
    "PERCENTILE.EXC", "QUARTILE.INC", "QUARTILE.EXC", "MODE.SNGL", "NORM.DIST", "NORM.S.DIST",
    "NORM.INV", "CEILING.MATH", "FLOOR.MATH", "FORECAST.LINEAR", "AGGREGATE",
}
NUMBER_FORMATS = {"integer": "#,##0", "number": "#,##0.00", "money": "#,##0.00", "percent": "0.0%"}
TOTAL_ROW = re.compile(r"^\s*(итого|итог|всего|total|subtotal|барлығы|жиыны|жалпы)", re.I)
_NUM_JUNK = re.compile(r"[\s']|₸|₽|€|\$|(?i:тг\.?|тенге|руб\.?|kzt|usd|eur|rub)")
_NUMBER = re.compile(r"[-+]?[\d.,]*\d[\d.,]*(?:[eE][-+]?\d+)?")

THIN = Side(style="thin", color="D9DCE6")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
ZEBRA_FILL = PatternFill("solid", fgColor="F7F6FF")
TOTAL_FILL = PatternFill("solid", fgColor="EDE9FE")
ANSWER_FILL = PatternFill("solid", fgColor="ECFDF5")


def xl_font(**style) -> Font:
    return Font(**{"name": "Calibri", "size": 11, **style})


def parse_number(text: str, kind: str = "number") -> int | float | None:
    """'1 250,5 ₸' → 1250.5, '15%' → 0.15; None, если это не число."""
    s = _NUM_JUNK.sub("", text or "").replace("−", "-").replace("–", "-")
    is_percent = s.endswith("%")
    s = s.rstrip("%")
    if not s or not _NUMBER.fullmatch(s):
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif s.count(",") == 1:
        s = s.replace(",", ".")
    s = s.replace(",", "")
    if s.count(".") > 1:
        s = s.replace(".", "")
    try:
        number = float(s)
    except ValueError:
        return None
    if not math.isfinite(number):
        return None
    if is_percent or (kind == "percent" and abs(number) > 1):
        return number / 100
    return int(number) if number.is_integer() and kind != "percent" else number


def _swap_separators(formula: str) -> str:
    """'=ROUND(A1;2)' → '=ROUND(A1,2)': точка с запятой вне строк и массивов."""
    out, quote, depth = [], None, 0
    for ch in formula:
        if quote:
            quote = None if ch == quote else quote
        elif ch in "\"'":
            quote = ch
        elif ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
        elif ch == ";" and depth == 0:
            ch = ","
        out.append(ch)
    return "".join(out)


def clean_formula(raw: str, renames: dict[str, str]) -> str | None:
    """Проверяет формулу от модели; None — если Excel её не поймёт (тогда пишем значение)."""
    formula = clean_text(raw)
    if not formula:
        return None
    if not formula.startswith("="):
        formula = "=" + formula
    if ";" in formula:
        if re.search(r"\d,\d", formula):  # локальный формат с десятичной запятой — не угадываем
            return None
        formula = _swap_separators(formula)
    for old, new in renames.items():
        formula = re.sub(rf"(?<![\w'])'?{re.escape(old)}'?!", f"'{new}'!", formula)
    try:
        tokenizer = Tokenizer(formula)
    except Exception:
        return None
    depth, prev_infix = 0, False
    for item in tokenizer.items:
        if item.type == Token.FUNC and item.subtype == Token.OPEN:
            name = item.value[:-1].upper().removeprefix("_XLFN.")
            if name in MODERN_FUNCTIONS:
                item.value = f"_xlfn.{name}("
            elif name in FORMULAE:
                item.value = f"{name}("
            else:
                return None
        if item.subtype == Token.OPEN:
            depth += 1
        elif item.subtype == Token.CLOSE:
            depth -= 1
            if depth < 0:
                return None
        infix = item.type == Token.OP_IN
        if infix and prev_infix:
            return None
        prev_infix = infix
    return tokenizer.render() if depth == 0 else None


def excel_value(cell: XlCell | None, kind: str, renames: dict[str, str]) -> Any:
    if cell is None:
        return None
    formula = clean_formula(cell.formula, renames) if cell.formula.strip() else None
    if formula:
        return formula
    text = clean_text(cell.value)
    if kind != "text":
        number = parse_number(text, kind)
        if number is not None:
            return number
    return text.lstrip("=") or None


def safe_sheet_title(name: str, used: set[str]) -> str:
    base = " ".join(re.sub(r"[\[\]:*?/\\]", " ", clean_text(name)).split()).strip("'")[:31] or "Лист"
    title, n = base, 2
    while title.lower() in used:
        suffix = f" ({n})"
        title, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(title.lower())
    return title


def text_row_height(text: str, chars_per_line: int) -> float:
    lines = sum(max(1, math.ceil(len(part) / chars_per_line)) for part in str(text).split("\n"))
    return max(15.0, 15.5 * lines)


def write_table_sheet(ws, sheet: XlSheet, renames: dict[str, str]) -> None:
    columns = sheet.columns
    widths = [min(len(clean_text(col.name)), 28) for col in columns]
    for c, col in enumerate(columns, 1):
        cell = ws.cell(row=1, column=c, value=clean_text(col.name) or f"Столбец {c}")
        cell.font = xl_font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.border = CELL_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32

    for r, row in enumerate(sheet.rows, 2):
        is_total = bool(row.cells) and bool(TOTAL_ROW.match(row.cells[0].value or ""))
        for c, col in enumerate(columns, 1):
            data = row.cells[c - 1] if c - 1 < len(row.cells) else None
            value = excel_value(data, col.kind, renames)
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = CELL_BORDER
            if col.kind in NUMBER_FORMATS:
                cell.number_format = NUMBER_FORMATS[col.kind]
            is_text = isinstance(value, str) and not value.startswith("=")
            cell.alignment = Alignment(vertical="center", wrap_text=is_text and len(value) > 45)
            if is_total:
                cell.font = xl_font(bold=True)
                cell.fill = TOTAL_FILL
                cell.border = Border(left=THIN, right=THIN, bottom=THIN, top=Side(style="medium", color="4F46E5"))
            elif r % 2 == 1:
                cell.fill = ZEBRA_FILL
            shown = data.value if data else ""
            widths[c - 1] = max(widths[c - 1], min(len(clean_text(shown)) + (4 if col.kind != "text" else 0), 50))

    for c, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = max(10, width + 3)
    ws.freeze_panes = "A2"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape" if len(columns) > 6 else "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def write_solution_sheet(ws, solution: ExcelSolution, labels: dict) -> None:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 60
    wrap_top = Alignment(wrap_text=True, vertical="top")
    accent = xl_font(bold=True, size=12, color="4338CA")

    ws["B1"] = clean_text(solution.title)
    ws["B1"].font = xl_font(bold=True, size=16, color="4338CA")
    ws["B2"] = f"{labels['created']} · {datetime.now():%d.%m.%Y}"
    ws["B2"].font = xl_font(italic=True, size=9, color="6B7280")

    def merged_text(row: int, text: str, **style) -> None:
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        cell = ws.cell(row=row, column=2, value=text)
        cell.alignment = wrap_top
        for key, value in style.items():
            setattr(cell, key, value)
        ws.row_dimensions[row].height = text_row_height(text, 120)

    row = 4
    ws.cell(row=row, column=2, value=labels["task"]).font = accent
    merged_text(row + 1, clean_text(solution.task))
    row += 3
    ws.cell(row=row, column=2, value=labels["steps"]).font = accent
    for i, step in enumerate(solution.steps, 1):
        row += 1
        number = ws.cell(row=row, column=1, value=f"{i}.")
        number.font = xl_font(bold=True, color="6D28D9")
        number.alignment = Alignment(horizontal="right", vertical="top")
        merged_text(row, clean_text(step).replace("**", ""))
    row += 2
    ws.cell(row=row, column=2, value=labels["answer"]).font = accent
    row += 1
    green = Side(style="thin", color="34D399")
    merged_text(row, clean_text(solution.answer).replace("**", ""), font=xl_font(bold=True, color="065F46"),
                fill=ANSWER_FILL)
    for column in (2, 3):
        ws.cell(row=row, column=column).border = Border(left=green, right=green, top=green, bottom=green)
        ws.cell(row=row, column=column).fill = ANSWER_FILL

    if solution.key_formulas:
        row += 2
        ws.cell(row=row, column=2, value=labels["formulas"]).font = accent
        row += 1
        for column, title in ((2, labels["formula"]), (3, labels["meaning"])):
            cell = ws.cell(row=row, column=column, value=title)
            cell.font, cell.fill, cell.border = xl_font(bold=True, color="FFFFFF"), HEADER_FILL, CELL_BORDER
        for note in solution.key_formulas:
            row += 1
            formula_cell = ws.cell(row=row, column=2, value=clean_text(note.formula))
            formula_cell.data_type = "s"  # показываем формулу как текст, а не вычисляем
            formula_cell.font = xl_font(name="Consolas", color="1E40AF")
            explanation = clean_text(note.explanation)
            meaning_cell = ws.cell(row=row, column=3, value=explanation)
            for cell in (formula_cell, meaning_cell):
                cell.border, cell.alignment = CELL_BORDER, wrap_top
            ws.row_dimensions[row].height = max(text_row_height(explanation, 62),
                                                text_row_height(note.formula, 72))
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def build_xlsx(solution: ExcelSolution, labels: dict) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    used = {labels["solution"].lower()}
    renames: dict[str, str] = {}
    titles = []
    for sheet in solution.sheets:
        title = safe_sheet_title(sheet.name, used)
        titles.append(title)
        if sheet.name and title != sheet.name:
            renames[sheet.name] = title
    for sheet, title in zip(solution.sheets, titles):
        if sheet.columns:
            write_table_sheet(workbook.create_sheet(title), sheet, renames)
    write_solution_sheet(workbook.create_sheet(labels["solution"]), solution, labels)
    for index, ws in enumerate(workbook.worksheets):
        ws.sheet_view.tabSelected = index == 0
    workbook.active = 0
    workbook.properties.title = clean_text(solution.title)
    workbook.properties.creator = APP_NAME
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def unique_names(names: Iterable[str]) -> list[str]:
    seen: dict[str, int] = {}
    result = []
    for i, name in enumerate(names, 1):
        name = clean_text(name) or f"Столбец {i}"
        seen[name] = seen.get(name, 0) + 1
        result.append(name if seen[name] == 1 else f"{name} ({seen[name]})")
    return result


def sheet_frame(sheet: XlSheet, show_formulas: bool) -> tuple[pd.DataFrame, dict]:
    """Таблица для предпросмотра: значения от AI или формулы."""
    names = unique_names(col.name for col in sheet.columns)
    data, config = {}, {}
    for idx, (name, col) in enumerate(zip(names, sheet.columns)):
        raw = []
        for row in sheet.rows:
            cell = row.cells[idx] if idx < len(row.cells) else None
            if cell is None:
                raw.append("")
            elif show_formulas and cell.formula.strip():
                raw.append(clean_formula(cell.formula, {}) or clean_text(cell.value))
            else:
                raw.append(clean_text(cell.value))
        numbers = [parse_number(v, col.kind) if v else None for v in raw]
        is_numeric = col.kind != "text" and all(n is not None for n, v in zip(numbers, raw) if v)
        if is_numeric and not (show_formulas and any(v.startswith("=") for v in raw)):
            data[name] = numbers
            config[name] = st.column_config.NumberColumn(
                format="percent" if col.kind == "percent" else "localized")
        else:
            data[name] = raw
    return pd.DataFrame(data), config


# ════════════════════════════════════════════════════════════════════════════
# Оформление (CSS) и HTML-компоненты
# ════════════════════════════════════════════════════════════════════════════

CSS = """
:root {
  --mk-bg: #0B0D12;
  --mk-surface: #121620;
  --mk-border: rgba(255, 255, 255, 0.07);
  --mk-border-strong: rgba(255, 255, 255, 0.13);
  --mk-muted: #98A1B3;
  --mk-violet: #A78BFA;
  --mk-grad: linear-gradient(135deg, #8B5CF6 0%, #6366F1 50%, #3B82F6 100%);
  --mk-radius: 16px;
}

/* Фон приложения с мягким свечением */
.stApp {
  background:
    radial-gradient(1100px 520px at 12% -8%, rgba(139, 92, 246, 0.13), transparent 60%),
    radial-gradient(900px 480px at 100% 0%, rgba(59, 130, 246, 0.09), transparent 55%),
    var(--mk-bg);
}
[data-testid="stHeader"] { background: transparent; }
[data-testid="stAppDeployButton"], [data-testid="stHeaderActionElements"] { display: none; }
[data-testid="stMainBlockContainer"] { max-width: 1120px; padding-top: 2.4rem; padding-bottom: 4rem; }
[data-testid="stBottomBlockContainer"] { max-width: 1120px; }
::-webkit-scrollbar { width: 10px; height: 10px; }
::-webkit-scrollbar-thumb { background: #262B38; border-radius: 10px; border: 2px solid transparent; background-clip: content-box; }

/* Сайдбар */
[data-testid="stSidebar"] { background: linear-gradient(180deg, #0F121A 0%, #0B0D12 100%); border-right: 1px solid var(--mk-border); }
[data-testid="stSidebarHeader"] { height: 1.5rem; min-height: 1.5rem; }
.mk-brand { display: flex; align-items: center; gap: 12px; margin: 4px 0 2px; }
.mk-logo { width: 42px; height: 42px; border-radius: 13px; background: var(--mk-grad); display: grid; place-items: center;
  font-weight: 800; font-size: 20px; color: #fff; box-shadow: 0 10px 26px -10px rgba(139, 92, 246, .8); }
.mk-brand-name { font-weight: 750; font-size: 1.06rem; letter-spacing: -0.01em; color: #fff; }
.mk-brand-sub { color: var(--mk-muted); font-size: .78rem; }
.mk-label { font-size: .68rem; font-weight: 650; letter-spacing: .09em; text-transform: uppercase; color: #6B7487; margin: 14px 0 0; }
.mk-foot { color: #5D6577; font-size: .74rem; line-height: 1.5; margin-top: 18px; padding-top: 14px; border-top: 1px solid var(--mk-border); }
.mk-link { color: var(--mk-violet) !important; text-decoration: none; font-size: .82rem; }

/* Статус подключения */
.mk-conn { display: flex; align-items: center; gap: 11px; padding: 11px 13px; border-radius: 13px;
  background: rgba(255, 255, 255, .03); border: 1px solid var(--mk-border); }
.mk-dot { width: 9px; height: 9px; border-radius: 50%; flex-shrink: 0; background: #6B7487; }
.mk-conn.ok .mk-dot { background: #34D399; box-shadow: 0 0 0 4px rgba(52, 211, 153, .16); }
.mk-conn.warn .mk-dot { background: #FBBF24; box-shadow: 0 0 0 4px rgba(251, 191, 36, .16); }
.mk-conn.err .mk-dot { background: #F87171; box-shadow: 0 0 0 4px rgba(248, 113, 113, .16); }
.mk-conn-title { font-weight: 650; font-size: .86rem; color: #fff; }
.mk-conn-sub { color: var(--mk-muted); font-size: .76rem; line-height: 1.4; word-break: break-word; }

/* Навигация по режимам (st.radio) */
.st-key-nav .stElementContainer, .st-key-nav [data-testid="stRadio"] { width: 100% !important; }
.st-key-nav [data-testid="stRadioGroup"] { gap: 4px; width: 100%; }
.st-key-nav [data-testid="stRadioGroup"] > div { width: 100%; padding: 9px 12px; border-radius: 12px;
  border: 1px solid transparent; transition: background .15s ease, border-color .15s ease; }
.st-key-nav [data-testid="stRadioGroup"] > div:hover { background: rgba(255, 255, 255, .04); }
.st-key-nav [data-testid="stRadioGroup"] > div[data-selected="true"] {
  background: linear-gradient(135deg, rgba(139, 92, 246, .22), rgba(59, 130, 246, .10));
  border-color: rgba(139, 92, 246, .38); }
.st-key-nav [data-testid="stRadioOption"] { width: 100%; cursor: pointer; }
.st-key-nav [data-testid="stRadioOption"] > div > div:first-child { display: none; }
.st-key-nav [data-testid="stRadioOption"] p { font-weight: 600; font-size: .93rem; }
.st-key-nav [data-testid="stRadioCaption"] { padding-left: 1.6rem; margin-left: 0; }
.st-key-nav [data-testid="stRadioCaption"] p { font-size: .76rem; }

/* Карточки (контейнеры с ключом card_*) */
[class*="st-key-card_"] { background: rgba(18, 22, 32, .72); border: 1px solid var(--mk-border) !important;
  border-radius: var(--mk-radius) !important; padding: 22px 22px 18px !important;
  box-shadow: 0 18px 40px -24px rgba(0, 0, 0, .8); }

/* Кнопки */
button[data-testid="stBaseButton-primary"] { background: var(--mk-grad); border: none; color: #fff; font-weight: 600;
  box-shadow: 0 10px 24px -12px rgba(124, 58, 237, .9); transition: transform .12s ease, filter .12s ease, box-shadow .12s ease; }
button[data-testid="stBaseButton-primary"]:hover { filter: brightness(1.1); transform: translateY(-1px); color: #fff;
  box-shadow: 0 14px 30px -12px rgba(124, 58, 237, 1); }
button[data-testid="stBaseButton-primary"]:active { transform: translateY(0); }
button[data-testid="stBaseButton-primary"]:disabled { background: #222736; color: #6B7487; box-shadow: none; transform: none; }
button[data-testid="stBaseButton-secondary"] { background: rgba(255, 255, 255, .03); border: 1px solid var(--mk-border-strong);
  transition: border-color .15s ease, background .15s ease; }
button[data-testid="stBaseButton-secondary"]:hover { border-color: rgba(139, 92, 246, .6); background: rgba(139, 92, 246, .08); color: #fff; }

/* Примеры-подсказки (карточки-кнопки одинаковой высоты в ряду) */
[class*="st-key-starters"] [data-testid="stColumn"] > [data-testid="stVerticalBlock"],
[class*="st-key-starters"] .stElementContainer, [class*="st-key-starters"] .stButton { height: 100%; }
[class*="st-key-starters"] button { min-height: 64px; height: 100%; justify-content: flex-start; padding: 12px 16px;
  border-radius: 14px; background: rgba(18, 22, 32, .7); }
[class*="st-key-starters"] button div, [class*="st-key-starters"] button span { justify-content: flex-start; text-align: left; }
[class*="st-key-starters"] button p { text-align: left; white-space: normal; line-height: 1.45; }
.st-key-xl_examples button { border-radius: 999px; padding: 4px 14px; min-height: 34px; }
.st-key-xl_examples button p { font-size: .85rem; }

/* Вкладки, экспандеры, чат */
[data-testid="stTabs"] [role="tablist"] { gap: 6px; border-bottom: 1px solid var(--mk-border); }
[data-testid="stTab"] { padding: 8px 12px; }
[data-testid="stTab"][aria-selected="true"] p { color: #fff; font-weight: 600; }
[data-testid="stExpander"] details { border: 1px solid var(--mk-border); border-radius: 14px; background: rgba(18, 22, 32, .6); }
[data-testid="stChatMessage"] { background: rgba(18, 22, 32, .7); border: 1px solid var(--mk-border); border-radius: 16px; padding: 14px 16px; }
[data-testid="stChatMessage"]:has([aria-label="Chat message from user"]) {
  background: linear-gradient(135deg, rgba(139, 92, 246, .13), rgba(59, 130, 246, .06)); border-color: rgba(139, 92, 246, .28); }
[data-testid="stChatMessage"]:has([aria-label="Chat message from assistant"]) [data-testid="stChatMessageAvatarCustom"] {
  background: var(--mk-grad); border: none; color: #fff; }
[data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] :is(h1, h2, h3) { font-size: 1.15rem; padding: .7rem 0 .3rem; }
[data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] :is(h4, h5, h6) { font-size: 1rem; }
[data-testid="stChatInput"] { border-radius: 16px; }
[data-testid="stBottom"] > div { background: linear-gradient(180deg, rgba(11, 13, 18, 0) 0%, var(--mk-bg) 32%); }

/* Шапка режима */
.mk-hero { position: relative; overflow: hidden; padding: 24px 26px; margin-bottom: 6px; border-radius: 20px;
  border: 1px solid var(--mk-border);
  background: linear-gradient(135deg, rgba(139, 92, 246, .15), rgba(59, 130, 246, .06) 55%, rgba(18, 22, 32, .5)); }
.mk-hero::after { content: ""; position: absolute; right: -70px; top: -90px; width: 260px; height: 260px;
  background: radial-gradient(circle, rgba(139, 92, 246, .35), transparent 70%); pointer-events: none; }
.mk-hero-row { display: flex; gap: 16px; align-items: center; position: relative; z-index: 1; }
.mk-hero-icon { width: 56px; height: 56px; border-radius: 17px; display: grid; place-items: center; font-size: 27px; flex-shrink: 0;
  background: rgba(255, 255, 255, .06); border: 1px solid var(--mk-border-strong); }
.mk-hero-title { font-size: 1.8rem; font-weight: 800; letter-spacing: -0.025em; line-height: 1.15;
  background: linear-gradient(90deg, #FFFFFF 0%, #C4B5FD 100%); -webkit-background-clip: text; background-clip: text; color: transparent; }
.mk-hero-sub { color: var(--mk-muted); margin-top: 6px; font-size: .95rem; line-height: 1.55; max-width: 760px; }
.mk-chips { display: flex; flex-wrap: wrap; gap: 8px; margin-top: 16px; position: relative; z-index: 1; }
.mk-chip { display: inline-flex; align-items: center; gap: 6px; padding: 4px 11px; border-radius: 999px; font-size: .76rem;
  color: #C9CEDA; background: rgba(255, 255, 255, .05); border: 1px solid var(--mk-border); }

/* Предупреждение об API-ключе */
.mk-banner { display: flex; gap: 12px; align-items: flex-start; padding: 14px 16px; border-radius: 14px; line-height: 1.55;
  background: rgba(251, 191, 36, .07); border: 1px solid rgba(251, 191, 36, .25); color: #FDE68A; font-size: .92rem; }
.mk-banner code { color: #FDE68A; background: rgba(251, 191, 36, .12); padding: 1px 6px; border-radius: 6px; }
.mk-banner a { color: #FCD34D; }

/* Результаты */
.mk-result { padding: 22px 24px; border-radius: 18px; border: 1px solid transparent;
  background: linear-gradient(#121620, #121620) padding-box, var(--mk-grad) border-box; }
.mk-eyebrow { font-size: .7rem; font-weight: 700; letter-spacing: .09em; text-transform: uppercase; color: var(--mk-violet); }
.mk-result-title { font-size: 1.45rem; font-weight: 750; letter-spacing: -0.015em; margin: 6px 0 8px; color: #fff; }
.mk-result-text { color: #C9CEDA; line-height: 1.65; }
.mk-answer { margin-top: 16px; padding: 14px 16px; border-radius: 12px; line-height: 1.6;
  background: rgba(52, 211, 153, .08); border: 1px solid rgba(52, 211, 153, .28); color: #D1FAE5; }
.mk-answer-label { font-size: .7rem; font-weight: 700; letter-spacing: .09em; text-transform: uppercase; color: #34D399; margin-bottom: 4px; }
.mk-ideas { display: grid; grid-template-columns: repeat(auto-fill, minmax(250px, 1fr)); gap: 12px; margin: 6px 0 14px; }
.mk-idea { padding: 16px 16px 14px; border-radius: 14px; background: var(--mk-surface); border: 1px solid var(--mk-border);
  transition: border-color .15s ease, transform .15s ease; }
.mk-idea:hover { border-color: rgba(139, 92, 246, .45); transform: translateY(-2px); }
.mk-idea-num { font-size: .72rem; font-weight: 800; color: var(--mk-violet); letter-spacing: .06em; }
.mk-idea-title { font-weight: 650; color: #fff; margin: 4px 0 6px; line-height: 1.35; }
.mk-idea-text { color: var(--mk-muted); font-size: .9rem; line-height: 1.55; }
.mk-section-title { font-size: 1rem; font-weight: 700; color: #fff; margin: 4px 0 6px; }
[class*="st-key-note_section_"] { border-radius: 14px !important; background: rgba(18, 22, 32, .55); }
[class*="st-key-note_section_"] h4 { font-size: 1.08rem; padding: .2rem 0 .4rem; color: #fff; }
.mk-callout { padding: 14px 16px; border-radius: 12px; background: rgba(139, 92, 246, .08);
  border-left: 3px solid #8B5CF6; color: #DDD6FE; line-height: 1.6; }
.mk-terms { display: grid; gap: 8px; }
.mk-term { display: grid; grid-template-columns: minmax(140px, 230px) 1fr; gap: 14px; padding: 12px 14px; border-radius: 12px;
  background: var(--mk-surface); border: 1px solid var(--mk-border); line-height: 1.55; }
.mk-term b { color: #C4B5FD; }
.mk-term span { color: #C9CEDA; }
.mk-steps { display: flex; flex-direction: column; gap: 10px; }
.mk-step { display: flex; gap: 12px; align-items: flex-start; padding: 12px 14px; border-radius: 12px;
  background: var(--mk-surface); border: 1px solid var(--mk-border); }
.mk-step-num { flex-shrink: 0; width: 26px; height: 26px; border-radius: 8px; background: var(--mk-grad); color: #fff;
  font-weight: 700; font-size: .8rem; display: grid; place-items: center; }
.mk-step-text { line-height: 1.6; color: #D5D9E3; padding-top: 2px; }
.mk-formula { padding: 12px 14px; border-radius: 12px; background: var(--mk-surface); border: 1px solid var(--mk-border); margin-bottom: 10px; }
.mk-formula code { color: #A5F3FC; background: rgba(34, 211, 238, .08); padding: 3px 8px; border-radius: 6px; font-size: .86rem; }
.mk-formula-text { color: var(--mk-muted); margin-top: 8px; font-size: .9rem; line-height: 1.55; }
.mk-files { display: flex; flex-wrap: wrap; gap: 6px; }
.mk-empty { text-align: center; padding: 26px 12px 6px; }
.mk-empty-icon { font-size: 38px; }
.mk-empty-title { font-size: 1.2rem; font-weight: 700; color: #fff; margin-top: 6px; }
.mk-empty-sub { color: var(--mk-muted); margin-top: 4px; font-size: .92rem; }
.mk-mini { font-size: .78rem; color: var(--mk-muted); margin: 2px 0 -4px; }

/* Индикаторы загрузки */
.mk-skeleton { padding: 22px; border-radius: 16px; background: var(--mk-surface); border: 1px solid var(--mk-border); }
.mk-sk { height: 12px; border-radius: 6px; margin: 12px 0;
  background: linear-gradient(90deg, rgba(255, 255, 255, .04) 25%, rgba(255, 255, 255, .11) 37%, rgba(255, 255, 255, .04) 63%);
  background-size: 400% 100%; animation: mk-shimmer 1.4s ease infinite; }
.mk-sk.title { height: 20px; width: 42%; margin-bottom: 20px; }
@keyframes mk-shimmer { 0% { background-position: 100% 50%; } 100% { background-position: 0 50%; } }
.mk-thinking { display: inline-flex; align-items: center; gap: 6px; color: var(--mk-muted); font-size: .9rem; padding: 4px 0; }
.mk-thinking i { width: 7px; height: 7px; border-radius: 50%; background: var(--mk-violet); display: inline-block;
  animation: mk-bounce 1.2s infinite ease-in-out; }
.mk-thinking i:nth-child(2) { animation-delay: .15s; }
.mk-thinking i:nth-child(3) { animation-delay: .3s; margin-right: 6px; }
@keyframes mk-bounce { 0%, 80%, 100% { transform: scale(.55); opacity: .4; } 40% { transform: scale(1); opacity: 1; } }

@media (max-width: 640px) {
  .mk-hero { padding: 18px; }
  .mk-hero-title { font-size: 1.35rem; }
  .mk-hero-icon { width: 46px; height: 46px; font-size: 22px; }
  .mk-term { grid-template-columns: 1fr; gap: 4px; }
}
"""

THINKING_HTML = '<div class="mk-thinking"><i></i><i></i><i></i>Gemini думает…</div>'


def skeleton_html() -> str:
    widths = (94, 81, 88, 67, 90, 74, 58)
    lines = "".join(f'<div class="mk-sk" style="width:{w}%"></div>' for w in widths)
    return f'<div class="mk-skeleton"><div class="mk-sk title"></div>{lines}</div>'


def chips_html(items: Iterable[str]) -> str:
    return '<div class="mk-chips">' + "".join(f'<span class="mk-chip">{esc(i)}</span>' for i in items) + "</div>"


def render_hero(mode_key: str, ctx: Ctx) -> None:
    mode = MODES[mode_key]
    chips = [f"✦ {ctx.model}", f"🌐 {ctx.lang_name}", f"🎓 {ctx.level}",
             "🧠 Глубокое мышление" if ctx.deep else "⚡ Быстрый режим"]
    st.html(
        f'<div class="mk-hero"><div class="mk-hero-row"><div class="mk-hero-icon">{mode.emoji}</div>'
        f'<div><div class="mk-hero-title">{esc(mode.title)}</div>'
        f'<div class="mk-hero-sub">{esc(mode.subtitle)}</div></div></div>{chips_html(chips)}</div>'
    )


def render_key_banner() -> None:
    st.html(
        '<div class="mk-banner"><span>🔑</span><div><b>Нужен API-ключ Gemini.</b> Добавьте строку '
        '<code>GEMINI_API_KEY=ваш_ключ</code> в файл <code>.env</code> рядом с app.py '
        f'или введите ключ в боковой панели. <a href="{API_KEY_URL}" target="_blank">Получить ключ бесплатно →</a>'
        "</div></div>"
    )


def progress_label(step: Callable[[str], None], text: str) -> Callable[[int], None]:
    """Колбэк прогресса: «Gemini пишет конспект · 2 340 симв.»."""
    return lambda chars: step(f"{text} · {chars:,} симв.".replace(",", " "))


def run_with_progress(label: str, done: str, model: str, work: Callable[[Callable[[str], None]], Any]) -> Any:
    """Выполняет work() со статусом, скелетоном загрузки и обработкой ошибок."""
    status = st.status(label, expanded=False)
    skeleton = st.empty()
    skeleton.html(skeleton_html())
    started = time.perf_counter()
    try:
        result = work(lambda text: status.update(label=text))
    except Exception as exc:
        skeleton.empty()
        status.update(label="Не получилось", state="error")
        st.error(describe_error(exc, model), icon=":material/error:")
        return None
    skeleton.empty()
    status.update(label=f"{done} за {time.perf_counter() - started:.1f} с", state="complete")
    return result


def set_state(key: str, value: Any) -> None:
    st.session_state[key] = value


def drop_state(*keys: str) -> None:
    for key in keys:
        st.session_state.pop(key, None)


# ════════════════════════════════════════════════════════════════════════════
# Чат (общий для Q&A по конспекту и режима ДЗ)
# ════════════════════════════════════════════════════════════════════════════

def to_user_message(value: Any) -> dict | None:
    text = (value if isinstance(value, str) else getattr(value, "text", "") or "").strip()
    files = [] if isinstance(value, str) else list(getattr(value, "files", None) or [])
    attachments = []
    for upload in files:
        ext = upload.name.rsplit(".", 1)[-1].lower()
        attachments.append({"name": upload.name, "mime": MIME_BY_EXT.get(ext, upload.type or "text/plain"),
                            "data": upload.getvalue()})
    if not text and not attachments:
        return None
    return {"role": "user", "text": text, "files": attachments}


def to_contents(messages: list[dict]) -> list[types.Content]:
    contents = []
    for msg in messages:
        parts = [types.Part.from_bytes(data=f["data"], mime_type=f["mime"]) for f in msg.get("files", [])]
        text = msg.get("text") or ("Разбери задание из вложения." if msg["role"] == "user" else "")
        if text:
            parts.append(types.Part.from_text(text=text))
        contents.append(types.Content(role="user" if msg["role"] == "user" else "model", parts=parts))
    return contents


def render_message(msg: dict) -> None:
    is_user = msg["role"] == "user"
    with st.chat_message("user" if is_user else "assistant", avatar=USER_AVATAR if is_user else AI_AVATAR):
        files = msg.get("files", [])
        images = [f["data"] for f in files if f["mime"].startswith("image/")]
        if images:
            st.image(images, width=180)
        others = [f["name"] for f in files if not f["mime"].startswith("image/")]
        if others:
            st.html('<div class="mk-files">' + "".join(f'<span class="mk-chip">📎 {esc(n)}</span>' for n in others)
                    + "</div>")
        if msg.get("text"):
            st.markdown(msg["text"])


def retry_failed(pending_key: str, failed_key: str) -> None:
    failed = st.session_state.pop(failed_key, None)
    if failed:
        st.session_state[pending_key] = failed["msg"]


def error_with_retry(message: str, pending_key: str, failed_key: str, button_key: str) -> None:
    st.error(message, icon=":material/error:")
    st.button("Повторить", icon=":material/refresh:", key=button_key, on_click=retry_failed,
              args=(pending_key, failed_key))


def with_indicator(indicator, chunks: Iterator[str]) -> Iterator[str]:
    """Показывает «Gemini думает…», пока не придёт первый фрагмент ответа."""
    try:
        first = True
        for chunk in chunks:
            if first:
                indicator.empty()
                first = False
            yield chunk
    finally:
        indicator.empty()


def answer_turn(ctx: Ctx, messages: list[dict], user_msg: dict, *, system: str, prefix: list,
                pending_key: str, failed_key: str) -> None:
    render_message(user_msg)
    with st.chat_message("assistant", avatar=AI_AVATAR):
        indicator = st.empty()
        indicator.html(THINKING_HTML)
        try:
            contents = prefix + to_contents(messages + [user_msg])
            answer = st.write_stream(with_indicator(indicator, stream_text(ctx, contents, system=system)))
            answer = answer if isinstance(answer, str) else "".join(map(str, answer))
            if not answer.strip():
                raise AppError("Gemini вернул пустой ответ. Попробуйте переформулировать вопрос.")
        except Exception as exc:
            indicator.empty()
            message = describe_error(exc, ctx.model)
            st.session_state[failed_key] = {"msg": user_msg, "error": message}
            error_with_retry(message, pending_key, failed_key, f"{failed_key}_now")
            return
    messages.extend([user_msg, {"role": "assistant", "text": answer}])


def render_starters(key: str, starters: list[tuple[str, str]], pending_key: str, disabled: bool) -> None:
    with st.container(key=f"starters_{key}"):
        for start in range(0, len(starters), 2):
            for column, (i, (emoji, prompt)) in zip(st.columns(2), enumerate(starters[start:start + 2], start)):
                column.button(
                    f"{emoji}  {prompt}", key=f"{key}_starter_{i}", width="stretch", disabled=disabled,
                    on_click=set_state, args=(pending_key, {"role": "user", "text": prompt, "files": []}),
                )


def chat_session(ctx: Ctx, *, key: str, system: str, placeholder: str, starters: list[tuple[str, str]],
                 prefix: list | None = None, accept_files: bool = False, empty_html: str = "") -> None:
    """Чат с историей, потоковыми ответами, подсказками и повтором при ошибке.

    Внутри контейнера (вкладки) поле ввода встроено в страницу, на верхнем уровне — закреплено внизу.
    """
    messages: list[dict] = st.session_state.setdefault(f"{key}_messages", [])
    pending_key, failed_key = f"{key}_pending", f"{key}_failed"

    history = st.container()
    starters_slot = st.empty()  # пустой слот сразу убирает подсказки прошлого запуска
    value = st.chat_input(
        placeholder if ctx.ready else "Сначала добавьте API-ключ в боковой панели",
        key=f"{key}_input",
        accept_file="multiple" if accept_files else False,
        file_type=CHAT_FILE_TYPES if accept_files else None,
        max_upload_size=MAX_FILE_MB if accept_files else None,
        disabled=not ctx.ready,
        submit_mode="disable",
    )
    user_msg = (to_user_message(value) if value else None) or st.session_state.pop(pending_key, None)
    if user_msg:
        st.session_state.pop(failed_key, None)
    failed = st.session_state.get(failed_key)

    with history:
        for msg in messages:
            render_message(msg)
        if user_msg:
            answer_turn(ctx, messages, user_msg, system=system, prefix=prefix or [],
                        pending_key=pending_key, failed_key=failed_key)
        elif failed:
            render_message(failed["msg"])
            with st.chat_message("assistant", avatar=AI_AVATAR):
                error_with_retry(failed["error"], pending_key, failed_key, f"{failed_key}_retry")

    if not messages and not user_msg and not failed:
        with starters_slot.container():
            if empty_html:
                st.html(empty_html)
            render_starters(key, starters, pending_key, disabled=not ctx.ready)


def chat_markdown(messages: list[dict]) -> str:
    lines = [f"# {APP_NAME} — диалог от {datetime.now():%d.%m.%Y %H:%M}", ""]
    for msg in messages:
        who = "Вы" if msg["role"] == "user" else APP_NAME
        files = ", ".join(f["name"] for f in msg.get("files", []))
        lines += [f"### {who}", *([f"*Вложения: {files}*"] if files else []), msg.get("text") or "", ""]
    return "\n".join(lines)


# ════════════════════════════════════════════════════════════════════════════
# Боковая панель
# ════════════════════════════════════════════════════════════════════════════

def connection_html(state: str, title: str, subtitle: str) -> str:
    return (f'<div class="mk-conn {state}"><span class="mk-dot"></span><div>'
            f'<div class="mk-conn-title">{esc(title)}</div><div class="mk-conn-sub">{esc(subtitle)}</div></div></div>')


def check_connection(api_key: str, model: str, slot) -> dict:
    """Проверяет ключ и модель (без расхода токенов). Результат кешируется в сессии."""
    fingerprint = hashlib.sha256(f"{api_key}|{model}".encode()).hexdigest()
    check = st.session_state.get("conn")
    if check and check["fp"] == fingerprint:
        return check
    slot.html(connection_html("warn", "Проверяю подключение…", model))
    try:
        info = get_client(api_key).models.get(
            model=model, config=types.GetModelConfig(http_options=types.HttpOptions(timeout=20_000)))
        check = {"fp": fingerprint, "ok": True, "title": "Gemini подключён", "detail": info.display_name or model}
    except Exception as exc:
        check = {"fp": fingerprint, "ok": False, "title": "Ошибка подключения", "detail": describe_error(exc, model)}
    st.session_state.conn = check
    return check


def save_env() -> None:
    api_key = st.session_state.get("api_key", "").strip()
    model = (st.session_state.get("model") or DEFAULT_MODEL).strip()
    try:
        ENV_PATH.touch(exist_ok=True)
        set_key(str(ENV_PATH), "GEMINI_API_KEY", api_key)
        set_key(str(ENV_PATH), "GEMINI_MODEL", model)
    except OSError as exc:
        st.session_state.flash = ("❌", f"Не удалось сохранить .env: {exc}")
        return
    os.environ["GEMINI_API_KEY"], os.environ["GEMINI_MODEL"] = api_key, model
    st.session_state.flash = ("✅", f"Ключ и модель сохранены в {ENV_PATH.name}")


def render_sidebar() -> tuple[str, Ctx]:
    with st.sidebar:
        st.html(
            '<div class="mk-brand"><div class="mk-logo">M</div><div>'
            f'<div class="mk-brand-name">{APP_NAME}</div>'
            '<div class="mk-brand-sub">Учебный ассистент на Gemini</div></div></div>'
        )

        st.html('<div class="mk-label">Подключение</div>')
        api_key = st.session_state.get("api_key", "").strip()
        model = (st.session_state.get("model") or DEFAULT_MODEL).strip()
        slot = st.empty()
        if not api_key:
            check = None
            slot.html(connection_html("warn", "Нужен API-ключ", "Укажите GEMINI_API_KEY в .env или ниже"))
        else:
            check = check_connection(api_key, model, slot)
            slot.html(connection_html("ok" if check["ok"] else "err", check["title"], check["detail"]))

        with st.expander("API-ключ и модель", icon=":material/key:", expanded=not (check and check["ok"])):
            st.text_input("GEMINI_API_KEY", key="api_key", type="password", placeholder="AIza…",
                          help="По умолчанию ключ читается из файла .env рядом с app.py.")
            st.selectbox("Модель Gemini", MODEL_OPTIONS, key="model", accept_new_options=True,
                         help="По умолчанию gemini-2.5-flash. Можно ввести название другой модели.")
            st.button("Проверить ключ", icon=":material/sync:", width="stretch", disabled=not api_key,
                      on_click=drop_state, args=("conn",))
            st.button("Сохранить в .env", icon=":material/save:", width="stretch", disabled=not api_key,
                      on_click=save_env, help="Записать ключ и модель в файл .env рядом с app.py")
            st.html(f'<a class="mk-link" href="{API_KEY_URL}" target="_blank">Получить бесплатный ключ →</a>')

        st.html('<div class="mk-label">Режим работы</div>')
        with st.container(key="nav"):
            mode = st.radio(
                "Режим работы", list(MODES), key="mode", label_visibility="collapsed",
                format_func=lambda m: f"{MODES[m].icon} {MODES[m].title}",
                captions=[m.caption for m in MODES.values()],
            )

        st.html('<div class="mk-label">Настройки ответов</div>')
        lang = st.selectbox("Язык ответов", list(LANGUAGES), key="language")
        level = st.selectbox("Уровень обучения", LEVELS, key="level")
        deep = st.toggle("Глубокое мышление", key="deep",
                         help="Gemini дольше «думает» перед ответом — точнее в задачах и расчётах. "
                              "Выключите для более быстрых ответов (Gemini 2.5 Flash).")
        st.html('<div class="mk-foot">Gemini API · Streamlit<br>AI может ошибаться — проверяйте важные '
                'факты и расчёты.</div>')

    client = get_client(api_key) if api_key else None
    return mode, Ctx(client=client, model=model, language=LANGUAGES[lang], lang_name=lang, level=level, deep=deep)


# ════════════════════════════════════════════════════════════════════════════
# Режим 1: Конспектирование и Q&A
# ════════════════════════════════════════════════════════════════════════════

def notes_payload(source: str, topic: str | None, text: str | None, upload) -> dict:
    """Готовит запрос к Gemini и контекст для последующих вопросов."""
    if source == "Тема":
        topic = (topic or "").strip()
        if len(topic) < 2:
            raise AppError("Введите тему конспекта.")
        return {"label": topic, "context": f"Тема: {topic}", "pdf": None,
                "contents": [f"Тема: {topic}\n\nСоставь учебный конспект по этой теме."]}
    if source == "Текст":
        text = (text or "").strip()
        if len(text) < 40:
            raise AppError("Текст слишком короткий — вставьте хотя бы абзац или выберите источник «Тема».")
        return {"label": shorten(text), "context": text, "pdf": None,
                "contents": [f'Учебный материал:\n"""\n{text}\n"""\n\nСоставь конспект по этому материалу.']}
    if upload is None:
        raise AppError("Загрузите файл с учебным материалом.")
    file_text, pdf = read_source_file(upload)
    if pdf is not None:
        return {"label": upload.name, "context": "", "pdf": pdf,
                "contents": [types.Part.from_bytes(data=pdf, mime_type="application/pdf"),
                             "Составь учебный конспект по этому документу."]}
    return {"label": upload.name, "context": file_text, "pdf": None,
            "contents": [f'Учебный материал из файла «{upload.name}»:\n"""\n{file_text}\n"""\n\n'
                         "Составь конспект по этому материалу."]}


def generate_notes(ctx: Ctx, payload: dict, detail: str, n_questions: int) -> None:
    labels = ctx.labels

    def work(step: Callable[[str], None]) -> tuple[StudyNotes, bytes]:
        notes = generate_json(ctx, payload["contents"], system=notes_system(ctx, detail, n_questions),
                              schema=StudyNotes, temperature=0.4,
                              on_progress=progress_label(step, "Gemini пишет конспект"))
        if not notes.sections and not notes.key_ideas:
            raise AppError("Gemini вернул пустой конспект. Уточните тему или добавьте текст.")
        step("Оформляю документ Word…")
        return notes, build_notes_docx(notes, labels, payload["label"])

    out = run_with_progress("Gemini изучает материал и составляет конспект…", "Конспект готов", ctx.model, work)
    if out is None:
        return
    notes, docx_bytes = out
    st.session_state.notes_result = {
        "notes": notes, "docx": docx_bytes, "markdown": notes_to_markdown(notes, labels),
        "file_stem": slugify(notes.title, "Конспект"), "label": payload["label"],
        "context": payload["context"], "pdf": payload["pdf"], "labels": labels,
    }
    drop_state("notes_chat_messages", "notes_chat_failed", "notes_chat_pending", "notes_tabs")


def notes_chat_prefix(result: dict) -> list[types.Content]:
    if result["pdf"] is None:
        return []
    return [
        types.Content(role="user", parts=[
            types.Part.from_bytes(data=result["pdf"], mime_type="application/pdf"),
            types.Part.from_text(text="Это исходный учебный материал для вопросов."),
        ]),
        types.Content(role="model", parts=[types.Part.from_text(text="Материал изучен. Готов отвечать на вопросы.")]),
    ]


def render_notes_result(ctx: Ctx, result: dict) -> None:
    notes: StudyNotes = result["notes"]
    labels = result["labels"]
    chips = [f"💡 Главных мыслей: {len(notes.key_ideas)}", f"📚 Разделов: {len(notes.sections)}",
             f"📖 Терминов: {len(notes.terms)}", f"❓ Вопросов: {len(notes.qa)}", f"📄 {shorten(result['label'], 40)}"]
    st.html(
        f'<div class="mk-result"><div class="mk-eyebrow">Конспект готов</div>'
        f'<div class="mk-result-title">{esc(notes.title)}</div>'
        f'<div class="mk-result-text">{md_inline_html(notes.summary)}</div>{chips_html(chips)}</div>'
    )
    word, markdown, reset = st.columns([1.3, 1, 1])
    word.download_button("Скачать Word (.docx)", data=result["docx"], file_name=f"{result['file_stem']}.docx",
                         mime=DOCX_MIME, type="primary", icon=":material/download:", on_click="ignore",
                         width="stretch", key="notes_docx")
    markdown.download_button("Markdown (.md)", data=result["markdown"].encode("utf-8"),
                             file_name=f"{result['file_stem']}.md", mime="text/markdown",
                             icon=":material/description:", on_click="ignore", width="stretch", key="notes_md")
    reset.button("Новый конспект", icon=":material/restart_alt:", width="stretch", key="notes_reset",
                 on_click=drop_state, args=("notes_result", "notes_chat_messages", "notes_chat_failed", "notes_tabs"))

    tab_notes, tab_terms, tab_quiz, tab_chat = st.tabs(
        ["📝 Конспект", "📖 Термины", "❓ Самопроверка", "💬 Вопросы по материалу"],
        key="notes_tabs", on_change="rerun",
    )
    with tab_notes:
        if notes.key_ideas:
            st.html(
                f'<div class="mk-section-title">{esc(labels["key_ideas"])}</div><div class="mk-ideas">'
                + "".join(
                    f'<div class="mk-idea"><div class="mk-idea-num">{i:02d}</div>'
                    f'<div class="mk-idea-title">{md_inline_html(idea.title)}</div>'
                    f'<div class="mk-idea-text">{md_inline_html(idea.text)}</div></div>'
                    for i, idea in enumerate(notes.key_ideas, 1)
                ) + "</div>"
            )
        for i, section in enumerate(notes.sections):
            with st.container(border=True, key=f"note_section_{i}"):
                st.markdown(f"#### {md_safe(section.heading)}\n"
                            + "\n".join(f"- {md_safe(point)}" for point in section.points))
        if notes.conclusion:
            st.html(f'<div class="mk-callout"><b>{esc(labels["conclusion"])}.</b> '
                    f"{md_inline_html(notes.conclusion)}</div>")
    with tab_terms:
        if notes.terms:
            st.html('<div class="mk-terms">' + "".join(
                f'<div class="mk-term"><b>{md_inline_html(t.term)}</b><span>{md_inline_html(t.definition)}</span></div>'
                for t in notes.terms) + "</div>")
        else:
            st.info("В этом конспекте нет отдельного списка терминов.")
    with tab_quiz:
        st.caption("Сначала ответьте сами, затем раскройте вопрос, чтобы проверить себя.")
        for i, item in enumerate(notes.qa, 1):
            with st.expander(f"**{i}.** {md_safe(item.question)}"):
                st.markdown(md_safe(item.answer))
    with tab_chat:
        st.caption("Задавайте вопросы по материалу — Gemini ответит с опорой на конспект и исходный текст.")
        chat_session(ctx, key="notes_chat", system=notes_chat_system(ctx, result), prefix=notes_chat_prefix(result),
                     placeholder="Спросите что-нибудь по материалу…", starters=NOTES_STARTERS)


def page_notes(ctx: Ctx) -> None:
    render_hero("notes", ctx)
    if not ctx.ready:
        render_key_banner()
    with st.container(key="card_notes_form"):
        source = st.segmented_control("Источник материала", NOTES_SOURCES, key="notes_source", required=True,
                                      format_func=lambda s: {"Тема": "🏷️ Тема", "Текст": "📋 Текст",
                                                             "Файл": "📎 Файл"}[s])
        topic = text = upload = None
        if source == "Тема":
            topic = st.text_input("Тема", key="notes_topic",
                                  placeholder="Например: Фотосинтез, Закон Ома, Великий шёлковый путь")
        elif source == "Текст":
            text = st.text_area("Текст лекции, параграфа или статьи", key="notes_text", height=230,
                                max_chars=MAX_TEXT_CHARS, placeholder="Вставьте сюда учебный текст…")
        else:
            upload = st.file_uploader(f"PDF, DOCX, TXT или MD — до {MAX_FILE_MB} МБ",
                                      type=["pdf", "docx", "txt", "md"], key="notes_file",
                                      max_upload_size=MAX_FILE_MB)
        left, right = st.columns(2)
        detail = left.select_slider("Детализация", options=list(NOTES_DETAIL), key="notes_detail")
        n_questions = right.slider("Вопросов для самопроверки", 3, 10, key="notes_nq")
        clicked = st.button("Создать конспект", type="primary", icon=":material/auto_awesome:", width="stretch",
                            disabled=not ctx.ready, key="notes_go")

    if clicked:
        try:
            payload = notes_payload(source, topic, text, upload)
        except AppError as exc:
            st.warning(str(exc), icon=":material/info:")
        else:
            generate_notes(ctx, payload, detail, n_questions)

    if result := st.session_state.get("notes_result"):
        render_notes_result(ctx, result)


# ════════════════════════════════════════════════════════════════════════════
# Режим 2: Генератор Excel-задач
# ════════════════════════════════════════════════════════════════════════════

def generate_excel(ctx: Ctx, task: str) -> None:
    labels = ctx.labels

    def work(step: Callable[[str], None]) -> tuple[ExcelSolution, bytes]:
        solution = generate_json(ctx, [f"Задача студента:\n{task}"], system=excel_system(ctx),
                                 schema=ExcelSolution, temperature=0.2,
                                 on_progress=progress_label(step, "Gemini решает задачу"))
        if not any(sheet.columns for sheet in solution.sheets):
            raise AppError("Gemini не вернул таблицу. Уточните условие задачи и попробуйте снова.")
        step("Собираю файл Excel с формулами…")
        return solution, build_xlsx(solution, labels)

    out = run_with_progress("Gemini решает задачу и строит таблицы…", "Решение готово", ctx.model, work)
    if out is None:
        return
    solution, xlsx_bytes = out
    st.session_state.excel_result = {"solution": solution, "xlsx": xlsx_bytes,
                                     "file_name": f"{slugify(solution.title, 'Задача')}.xlsx"}
    drop_state("excel_tabs")


def render_excel_result(result: dict) -> None:
    solution: ExcelSolution = result["solution"]
    n_rows = sum(len(sheet.rows) for sheet in solution.sheets)
    n_formulas = sum(1 for sheet in solution.sheets for row in sheet.rows for cell in row.cells if cell.formula.strip())
    chips = [f"📄 Листов: {len(solution.sheets) + 1}", f"🧾 Строк данных: {n_rows}", f"ƒx Формул: {n_formulas}"]
    st.html(
        f'<div class="mk-result"><div class="mk-eyebrow">Решение задачи</div>'
        f'<div class="mk-result-title">{esc(solution.title)}</div>'
        f'<div class="mk-result-text">{md_inline_html(solution.task)}</div>'
        f'<div class="mk-answer"><div class="mk-answer-label">Ответ</div>{md_inline_html(solution.answer)}</div>'
        f"{chips_html(chips)}</div>"
    )
    download, reset = st.columns([1.4, 1])
    download.download_button("Скачать Excel (.xlsx)", data=result["xlsx"], file_name=result["file_name"],
                             mime=XLSX_MIME, type="primary", icon=":material/download:", on_click="ignore",
                             width="stretch", key="excel_download")
    reset.button("Новая задача", icon=":material/restart_alt:", width="stretch", key="excel_reset",
                 on_click=drop_state, args=("excel_result", "excel_tabs"))

    tab_steps, tab_tables, tab_formulas = st.tabs(["🧮 Решение", "📋 Таблицы", "ƒx Формулы"],
                                                  key="excel_tabs", on_change="rerun")
    with tab_steps:
        st.html('<div class="mk-steps">' + "".join(
            f'<div class="mk-step"><div class="mk-step-num">{i}</div>'
            f'<div class="mk-step-text">{md_inline_html(step)}</div></div>'
            for i, step in enumerate(solution.steps, 1)) + "</div>")
    with tab_tables:
        show_formulas = st.toggle("Показать формулы вместо значений", key="xl_show_formulas")
        for sheet in solution.sheets:
            if not sheet.columns:
                continue
            st.markdown(f"**{md_safe(sheet.name) or 'Лист'}**")
            frame, config = sheet_frame(sheet, show_formulas)
            st.dataframe(frame, hide_index=True, width="stretch", column_config=config, placeholder="",
                         height=min(38 + 35 * max(len(frame), 1), 560))
        st.caption("В предпросмотре — значения, рассчитанные AI. В файле Excel формулы пересчитываются "
                   "автоматически, так что можно менять исходные данные.")
    with tab_formulas:
        if solution.key_formulas:
            st.html("".join(
                f'<div class="mk-formula"><code>{esc(note.formula)}</code>'
                f'<div class="mk-formula-text">{md_inline_html(note.explanation)}</div></div>'
                for note in solution.key_formulas))
        else:
            st.info("В этом решении нет отдельного списка формул.")


def page_excel(ctx: Ctx) -> None:
    render_hero("excel", ctx)
    if not ctx.ready:
        render_key_banner()
    with st.container(key="card_excel_form"):
        task = st.text_area(
            "Условие задачи", key="xl_task", height=170, max_chars=20_000,
            placeholder="Опишите задачу: исходные данные и что нужно найти. Например: рассчитай амортизацию "
                        "оборудования стоимостью 2 000 000 ₸ линейным способом на 5 лет…",
        )
        st.html('<div class="mk-mini">Или выберите пример:</div>')
        with st.container(horizontal=True, gap="small", key="xl_examples"):
            for i, (label, example) in enumerate(EXCEL_EXAMPLES):
                st.button(label, key=f"xl_example_{i}", on_click=set_state, args=("xl_task", example))
        clicked = st.button("Решить и создать Excel", type="primary", icon=":material/bolt:", width="stretch",
                            disabled=not ctx.ready, key="xl_go")

    if clicked:
        if len(task.strip()) < 10:
            st.warning("Опишите задачу подробнее: какие есть данные и что нужно найти.", icon=":material/info:")
        else:
            generate_excel(ctx, task.strip())

    if result := st.session_state.get("excel_result"):
        render_excel_result(result)


# ════════════════════════════════════════════════════════════════════════════
# Режим 3: Быстрое решение ДЗ
# ════════════════════════════════════════════════════════════════════════════

def page_homework(ctx: Ctx) -> None:
    render_hero("homework", ctx)
    if not ctx.ready:
        render_key_banner()
    subject_col, style_col, save_col, new_col = st.columns([1.5, 1.5, 0.9, 0.9], vertical_alignment="bottom")
    subject = subject_col.selectbox("Предмет", SUBJECTS, key="hw_subject")
    style = style_col.segmented_control("Формат ответа", list(HW_STYLES), key="hw_style", required=True)
    new_col.button("Новый чат", icon=":material/add_comment:", width="stretch", key="hw_new",
                   on_click=drop_state, args=("hw_messages", "hw_failed", "hw_pending"))

    empty_html = ('<div class="mk-empty"><div class="mk-empty-icon">✨</div>'
                  '<div class="mk-empty-title">Чем помочь с учёбой?</div>'
                  '<div class="mk-empty-sub">Напишите вопрос внизу или прикрепите фото задачи (📎). '
                  "Можно начать с примера:</div></div>")
    chat_session(ctx, key="hw", system=homework_system(ctx, subject, style or "Подробно"),
                 placeholder="Задайте вопрос или прикрепите фото задачи…", starters=HW_STARTERS,
                 accept_files=True, empty_html=empty_html)

    # Кнопка сохранения рисуется после чата, чтобы учесть только что полученный ответ.
    messages = st.session_state.get("hw_messages", [])
    save_col.download_button("Сохранить", data=chat_markdown(messages).encode("utf-8"),
                             file_name=f"dialog_{datetime.now():%Y%m%d_%H%M}.md", mime="text/markdown",
                             icon=":material/download:", on_click="ignore", width="stretch",
                             disabled=not messages, key="hw_save", help="Скачать диалог в Markdown")


# ════════════════════════════════════════════════════════════════════════════
# Точка входа
# ════════════════════════════════════════════════════════════════════════════

STATE_DEFAULTS = {
    "mode": "notes", "language": "Русский", "level": "Университет", "deep": True, "model": DEFAULT_MODEL,
    "notes_source": "Тема", "notes_topic": "", "notes_text": "", "notes_detail": "Стандартно", "notes_nq": 5,
    "xl_task": "", "hw_subject": SUBJECTS[0], "hw_style": "Подробно",
}


def init_state() -> None:
    if "api_key" not in st.session_state:
        st.session_state.api_key = os.getenv("GEMINI_API_KEY", "").strip()
    for key, value in STATE_DEFAULTS.items():
        # Повторное присваивание сохраняет значения виджетов скрытых режимов между переключениями.
        st.session_state[key] = st.session_state.get(key, value)


def main() -> None:
    st.set_page_config(page_title=APP_NAME, page_icon="🎓", layout="wide", initial_sidebar_state="expanded")
    st.html(f"<style>{CSS}</style>")
    init_state()
    if flash := st.session_state.pop("flash", None):
        st.toast(flash[1], icon=flash[0])
    mode, ctx = render_sidebar()
    {"notes": page_notes, "excel": page_excel, "homework": page_homework}[mode](ctx)


if __name__ == "__main__":
    if st.runtime.exists():
        main()
    else:
        # Запуск через `python app.py`: перезапускаем через Streamlit с тёмной темой.
        from streamlit.web import cli as stcli

        sys.argv = ["streamlit", "run", str(Path(__file__).resolve()), "--theme.base", "dark"]
        sys.exit(stcli.main())
